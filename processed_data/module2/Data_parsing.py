import os
import torch
import re
import unicodedata
import pdfplumber
from pptx import Presentation
from transformers import AutoModel, AutoTokenizer
from llama_index.embeddings.huggingface import HuggingFaceEmbedding
from llama_index.core import VectorStoreIndex, SimpleDirectoryReader, Settings, get_response_synthesizer
from llama_index.core.ingestion import IngestionPipeline
from llama_index.core.node_parser import SentenceSplitter
from llama_index.vector_stores.chroma import ChromaVectorStore
from chromadb import PersistentClient
from llama_index.core.retrievers import BaseRetriever
from llama_index.core.query_engine import CustomQueryEngine
from llama_index.core.response_synthesizers import BaseSynthesizer
from llama_index.llms.ollama import Ollama
from llama_index.core.readers.base import BaseReader
from llama_index.core.schema import Document

# ------------------------------
# CUSTOM PPTX READER
# ------------------------------
class PPTXTextOnlyReader(BaseReader):
    def load_data(self, file_path: str, extra_info=None) -> list[Document]:
        prs = Presentation(file_path)
        text_runs = []

        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        cleaned_text = preprocess_text(file_path, text)  # Apply preprocessing
                        text_runs.append(cleaned_text)

        full_text = "\n".join(text_runs)
        return [Document(text=full_text, metadata={"file_path": str(file_path)})]

# ------------------------------
# CUSTOM PDF READER
# ------------------------------
class PDFTextOnlyReader(BaseReader):
    def load_data(self, file_path: str, extra_info=None) -> list[Document]:
        text_runs = []

        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    cleaned_text = preprocess_text(file_path, text)  # Apply preprocessing
                    text_runs.append(cleaned_text)

        full_text = "\n".join(text_runs)
        return [Document(text=full_text, metadata={"file_path": str(file_path)})]

# ------------------------------
# MODEL & LLM SETTINGS
# ------------------------------
try:
    model_name = "sentence-transformers/all-MiniLM-L6-v2"
    Settings.embed_model = HuggingFaceEmbedding(model_name=model_name)
    print(f"✅ Model {model_name} loaded successfully.")
except Exception as e:
    print(f"❌ Error loading embedding model: {e}")
    exit()

Settings.llm = Ollama(model="llama3.1:latest", request_timeout=120.0)

# ------------------------------
# TEXT PREPROCESSING FUNCTION
# ------------------------------
def preprocess_text(file_path, text):
    code_extensions = {".py", ".java", ".cpp", ".js", ".c", ".cs", ".html", ".css", ".php", ".rb"}
    text_extensions = {".pdf", ".docx", ".pptx", ".txt"}

    ext = os.path.splitext(file_path)[-1].lower()

    if ext in text_extensions:
        text = clean_text(text)
    return text

def clean_text(text):
    text = unicodedata.normalize("NFKD", text)

    # Preserve email addresses
    email_pattern = r'([A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,})'
    emails = re.findall(email_pattern, text)
    for i, email in enumerate(emails):
        text = text.replace(email, f'EMAIL_PLACEHOLDER_{i}')

    # Preserve URLs
    url_pattern = r'http[s]?://(?:[a-zA-Z]|[0-9]|[$-_@.&+]|[!*\\(\\),]|(?:%[0-9a-fA-F][0-9a-fA-F]))+'
    urls = re.findall(url_pattern, text)
    for i, url in enumerate(urls):
        text = text.replace(url, f'URL_PLACEHOLDER_{i}')

    # Remove table of contents artifacts (long sequences of dots)
    text = re.sub(r'\.{5,}', ' ', text)  # Remove sequences of 5 or more dots

    # Remove excessive spaces and unnecessary formatting
    text = re.sub(r'\s+', ' ', text).strip()  

    # Restore emails and URLs
    for i, email in enumerate(emails):
        text = text.replace(f'EMAIL_PLACEHOLDER_{i}', email)
    for i, url in enumerate(urls):
        text = text.replace(f'URL_PLACEHOLDER_{i}', url)

    return text



# ------------------------------
# LOAD DOCUMENTS & PREPROCESS
# ------------------------------
doc_path = "/Users/liteshperumalla/Desktop/Files/masters/Smart AI Tutor/Modules/"

try:
    reader = SimpleDirectoryReader(
        input_dir=doc_path,
        required_exts=['.pptx', '.ipynb', '.docx', '.csv', '.jpeg', '.pdf', '.png', '.py'],
        file_extractor={".pptx": PPTXTextOnlyReader(), ".pdf": PDFTextOnlyReader()}, 
        recursive=True
    )
    docs = reader.load_data()

    if not docs:
        print("❌ No documents found! Check the path and file extensions.")
        exit()

    print(f"✅ Loaded {len(docs)} docs")

except Exception as e:
    print(f"❌ Error loading documents: {e}")
    exit()

# ------------------------------
# PROCESS DOCUMENTS & CHUNKING
# ------------------------------
processed_docs = []
CHUNK_SIZE = 512
CHUNK_OVERLAP = 50

def chunk_text(text, chunk_size=512, overlap=50):
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end]
        chunks.append(chunk)
        start += chunk_size - overlap
    return chunks

for doc in docs:
    doc_text = preprocess_text(doc.metadata.get("file_path", ""), doc.get_content())
    text_chunks = chunk_text(doc_text)

    folder_name = os.path.basename(os.path.dirname(doc.metadata.get("file_path", "")))
    for chunk in text_chunks:
        metadata = {
            "file_name": str(doc.metadata.get("file_name", "")),
            "file_path": str(doc.metadata.get("file_path", "")),
            "folder_name": str(folder_name),
            "num_tokens": int(len(chunk.split())),
            "num_chars": int(len(chunk)),
        }
        processed_docs.append({
            "doc_id": doc.doc_id,
            "text": chunk,
            "metadata": metadata,
            "category": "<category>"
        })
# ------------------------------
# PROCESS DOCUMENTS
# ------------------------------
processed_docs = []
for doc in docs:
    text_chunks = chunk_text(doc.get_content())
    folder_name = os.path.basename(os.path.dirname(doc.metadata.get("file_path", "")))

    for chunk in text_chunks:
        metadata = {
            "file_name": str(doc.metadata.get("file_name", "")),
            "file_path": str(doc.metadata.get("file_path", "")),  # Ensure string type
            "folder_name": str(folder_name),
            "num_tokens": int(len(chunk.split())),
            "num_chars": int(len(chunk)),
        }
        processed_docs.append({
            "doc_id": doc.doc_id,
            "text": chunk,
            "metadata": metadata,
            "category": "<category>"
        })

# Print processed chunks
for i, doc in enumerate(processed_docs):
    print(f"\nChunk {i + 1}:")
    print(f"Document ID: {doc['doc_id']}")
    print(f"Metadata: {doc['metadata']}")
    print(f"Text: {doc['text']}\n")


# ------------------------------
# INITIALIZE CHROMADB VECTOR STORE
# ------------------------------
chroma_path = "./chroma_db"
try:
    chroma_client = PersistentClient(path=chroma_path)
    collection = chroma_client.get_or_create_collection("document_chunks")

    vector_store = ChromaVectorStore(chroma_client, collection_name="document_chunks")
    print("✅ ChromaDB initialized successfully.")
except Exception as e:
    print(f"❌ Error initializing ChromaDB: {e}")
    exit()


# ------------------------------
# DOCUMENT PROCESSING PIPELINE
# ------------------------------
pipeline = IngestionPipeline(
    transformations=[
        SentenceSplitter(chunk_size=128, chunk_overlap=10),
        Settings.embed_model
    ],
)

try:
    nodes = pipeline.run(documents=docs)
    if not nodes:
        print("❌ No nodes were created. Check document parsing.")
        exit()
    print(f"✅ {len(nodes)} document nodes created and stored in ChromaDB.")

    for i, node in enumerate(nodes):
        collection.add(
            ids=[str(i)],
            documents=[node.text],
            metadatas=[node.metadata]
        )
except Exception as e:
    print(f"❌ Error during ingestion pipeline: {e}")
    exit()


# ------------------------------
# CREATE VECTOR STORE INDEX
# ------------------------------
try:
    index = VectorStoreIndex(nodes, vector_store=vector_store)
    print("✅ Vector store index created successfully.")
    persist_dir = "./persisted_index"
    os.makedirs(persist_dir, exist_ok=True)
    index.storage_context.persist(persist_dir=persist_dir)
    print(f"✅ Index persisted to {persist_dir}")
except Exception as e:
    print(f"❌ Error creating VectorStoreIndex: {e}")
    exit()


# ------------------------------
# CREATE CHAT ENGINE & PROCESS QUERY
# ------------------------------
class RAGQueryEngine(CustomQueryEngine):
    """RAG Query Engine for custom retrieval and response synthesis."""

    retriever: BaseRetriever
    response_synthesizer: BaseSynthesizer

    def custom_query(self, query_str: str):
        nodes = self.retriever.retrieve(query_str)
        response_obj = self.response_synthesizer.synthesize(query_str, nodes)
        return response_obj


retriever = index.as_retriever()
synthesizer = get_response_synthesizer(response_mode="compact")

query_engine = RAGQueryEngine(
    retriever=retriever, response_synthesizer=synthesizer
)

response = query_engine.query("Write a code to find a factorial for a number?")
print(response)
